"""Read CSV and Excel ticket exports."""
from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class FileReadError(ValueError):
    pass


def read_tabular_file(file_obj, filename: str) -> tuple[list[str], list[dict[str, Any]]]:
    extension = Path(filename).suffix.lower()
    if extension not in {".csv", ".xlsx", ".xlsm"}:
        raise FileReadError("Unsupported file type. Use CSV, XLSX or XLSM.")
    try:
        if extension == ".csv":
            payload = file_obj.read()
            text = payload.decode("utf-8-sig") if isinstance(payload, bytes) else payload
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            matrix = list(csv.reader(text.splitlines(), dialect))
        else:
            workbook = load_workbook(file_obj, read_only=True, data_only=True)
            worksheet = workbook.active
            matrix = list(worksheet.iter_rows(values_only=True))
            workbook.close()
    except (OSError, UnicodeDecodeError, ValueError, csv.Error) as exc:
        raise FileReadError(f"Unable to read uploaded file: {exc}") from exc
    if not matrix:
        raise FileReadError("The uploaded file is empty.")
    headers = [str(v).strip() if v is not None else "" for v in matrix[0]]
    nonempty = [h.casefold() for h in headers if h]
    if not any(nonempty):
        raise FileReadError("The first row must contain column headers.")
    if len(nonempty) != len(set(nonempty)):
        raise FileReadError("Duplicate column headers are not allowed.")
    rows = []
    for number, raw in enumerate(matrix[1:], start=2):
        values = list(raw) + [None] * max(0, len(headers) - len(raw))
        if not any(v not in (None, "") for v in values):
            continue
        rows.append({"__row_number__": number, **{
            header: values[i] for i, header in enumerate(headers) if header
        }})
    return headers, rows
